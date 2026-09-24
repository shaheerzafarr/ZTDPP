from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/work/ZTDPP_Defense_Guide_and_Project_Pitch.docx")

NAVY = "10263F"
BLUE = "2F73B9"
CORAL = "E46F51"
GREEN = "5F8F78"
AMBER = "D69C45"
PAPER = "F6F2E9"
MIST = "E6EDF2"
PALE_BLUE = "DDEAF5"
PALE_GREEN = "DDE9E1"
PALE_AMBER = "F4E8CE"
PALE_RED = "F5DFD9"
WHITE = "FFFFFF"
SLATE = "566574"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color="CAD3DA", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = tc_borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def table(doc, headers, rows, widths, header_fill=NAVY):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    tbl.alignment = 1
    tbl_pr = tbl._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is not None:
        tbl_w.set(qn("w:w"), "9360")
        tbl_w.set(qn("w:type"), "dxa")
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = tbl.rows[0].cells[index]
        set_cell_width(cell, width)
        shade(cell, header_fill)
        borders(cell, "FFFFFF", "4")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    set_repeat_header(tbl.rows[0])
    for row_idx, values in enumerate(rows):
        cells = tbl.add_row().cells
        prevent_split(tbl.rows[-1])
        for index, (value, width) in enumerate(zip(values, widths)):
            set_cell_width(cells[index], width)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cells[index], WHITE if row_idx % 2 == 0 else "EEF2F5")
            borders(cells[index])
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(NAVY)
    return tbl


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def page_title(doc, number, heading, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{number:02d}  ")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(CORAL)
    r = p.add_run(heading)
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.style = doc.styles["Subtitle"]


def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullets(doc, items, color=BLUE):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(lead)
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(color)
            p.add_run(rest)
        else:
            p.add_run(item)


def numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{index}. ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(CORAL)
        p.add_run(item)


def callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    set_cell_width(tbl.cell(0, 0), 1.25)
    set_cell_width(tbl.cell(0, 1), 5.25)
    shade(tbl.cell(0, 0), accent)
    shade(tbl.cell(0, 1), fill)
    for cell in tbl.rows[0].cells:
        borders(cell, fill, "0")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = tbl.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label.upper())
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(255, 255, 255)
    p = tbl.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def script_box(doc, speaker, text):
    callout(doc, f"{speaker} says", f'“{text}”', PALE_GREEN, GREEN)


def new_page(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string("25384A")
styles["Normal"].paragraph_format.line_spacing = 1.08
styles["Title"].font.name = "Georgia"
styles["Title"].font.size = Pt(34)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor.from_string(NAVY)
styles["Subtitle"].font.name = "Aptos"
styles["Subtitle"].font.size = Pt(14)
styles["Subtitle"].font.italic = True
styles["Subtitle"].font.color.rgb = RGBColor.from_string(SLATE)
for name, size, color in [("Heading 1", 22, NAVY), ("Heading 2", 15, BLUE), ("Heading 3", 12, CORAL)]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string(color)
    styles[name].paragraph_format.space_before = Pt(8)
    styles[name].paragraph_format.space_after = Pt(5)

header = section.header.paragraphs[0]
header.text = "ZTDPP  |  DEFENSE GUIDE AND PROJECT PITCH"
header.style = styles["Caption"]
header.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
header.runs[0].font.bold = True
footer = section.footer.paragraphs[0]
footer.add_run("FAST-NUCES Karachi  •  Final Year Project  •  21 September 2026")
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
add_page_number(footer)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(62)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ZERO TRUST DIGITAL\nPROVENANCE PLATFORM")
r.bold = True
r.font.size = Pt(17)
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
p.add_run("Defense Guide &\nProject Pitch")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
r = p.add_run("A simple explanation of what we built, how it works, and how to present it confidently")
r.italic = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor.from_string(SLATE)

doc.add_paragraph().paragraph_format.space_after = Pt(28)
table(doc, ["Presenter", "Primary responsibility"], [
    ["Hasham", "Frontend experience and live user journey"],
    ["Sohaib", "Backend, security, provenance, ledger, and trust orchestration"],
    ["Shaheer", "AI detection model, training, inference service, and model limitations"],
], [1.65, 4.85])
doc.add_paragraph().paragraph_format.space_after = Pt(20)
callout(doc, "Core message", "ZTDPP does not rely on one signal. It combines signed provenance, tamper-evident history, metadata, perceptual matching, and AI evidence to explain whether digital content can be trusted.", PALE_BLUE, BLUE)

# 2
new_page(doc)
page_title(doc, 1, "The 60-Second Project Pitch", "Start with the problem, then explain the value in plain language.")
doc.add_heading("The pitch", level=2)
script_box(doc, "Hasham", "Today, digital images can be copied, edited, stripped of metadata, or generated by AI in seconds. A viewer normally sees only the final file and has no reliable way to know where it came from or what happened to it.")
script_box(doc, "Sohaib", "ZTDPP creates a verifiable history for digital content. When trusted content is registered, the platform records hashes, metadata, ownership, a signed manifest, and a tamper-evident ledger receipt. When an image is checked later, the system recomputes the evidence instead of trusting stored claims blindly.")
script_box(doc, "Shaheer", "AI detection adds another signal by estimating whether the pixels look AI-generated. We combine that result with provenance and metadata evidence to produce a clear verdict and an explainable report.")
doc.add_heading("One sentence to remember", level=2)
callout(doc, "Thesis", "ZTDPP is an evidence-combination platform: it proves trusted history when that history exists and uses forensic signals when it does not.", PALE_AMBER, AMBER)
doc.add_heading("What makes the project valuable", level=2)
bullets(doc, [
    ("Explainable: ", "the result shows the evidence behind the verdict."),
    ("Tamper-aware: ", "hashes, signatures, lineage, and ledger links are recomputed."),
    ("Practical: ", "users can verify images through a web dashboard or an API."),
    ("Extensible: ", "the AI model is a separate service and can be upgraded independently."),
])

# 3
new_page(doc)
page_title(doc, 2, "The Problem and Our Answer", "Explain the need without using heavy technical language.")
table(doc, ["Real-world problem", "ZTDPP answer"], [
    ["A file can be edited after it leaves its creator.", "Versioned manifests, parent-child lineage, and revocation preserve the history."],
    ["Metadata can be removed or changed.", "The platform treats metadata as evidence, then cross-checks hashes, signatures, and ledger records."],
    ["The same image may be resized or recompressed.", "Perceptual hashes help find visually similar registered content."],
    ["AI-generated media may have no trusted origin record.", "The AI service provides a pixel-based probability as a supporting signal."],
    ["A single score can be misleading.", "The trust engine returns the score breakdown, signals, matched claim, and verification checks."],
], [3.05, 3.45])
doc.add_heading("Our key design decision", level=2)
body(doc, "We do not say that AI detection alone proves whether an image is real or fake. AI models can fail on new generators or edited files. ZTDPP therefore combines independent evidence and clearly reports when evidence is missing, conflicting, or unavailable.")
callout(doc, "Defense point", "If a reviewer asks, “Why not just use the AI model?”, answer: because provenance proves history, while AI only estimates from visual patterns. The two signals solve different parts of the problem.", PALE_RED, CORAL)

# 4
new_page(doc)
page_title(doc, 3, "How One Image Moves Through ZTDPP", "Use this page as the simplest end-to-end explanation.")
numbered(doc, [
    "A creator or authorized user registers an image or a privacy-preserving content hash.",
    "The backend validates the request and extracts SHA-256, perceptual hashes, and available metadata.",
    "The platform creates a canonical manifest describing the content, source, action, ownership, and evidence.",
    "The manifest is signed with the platform Ed25519 signing key and added to the provenance ledger.",
    "If the image is edited, the new version points to its parent so the full lineage stays visible.",
    "During verification, the image remains in memory while hashes and metadata are extracted.",
    "The backend searches for an exact registered claim, then tries a perceptual near match.",
    "At the same time, the backend calls the FastAPI AI service for a pixel-based prediction.",
    "The trust engine combines the available evidence, applies safety caps, and produces a verdict.",
    "The complete report is saved so the result can be reviewed later with the same evidence and policy snapshot.",
])
doc.add_heading("What is returned to the user?", level=2)
callout(doc, "Response", "Verdict, classification, trust score, metadata score, AI score, effective weights, evidence signals, matched provenance, signature checks, ledger checks, extracted metadata, and processing time.", PALE_BLUE, BLUE)

# 5
new_page(doc)
page_title(doc, 4, "Architecture in Easy Language", "Each service has one clear responsibility.")
table(doc, ["Layer", "Technology", "Simple responsibility", "Owner"], [
    ["Web application", "Next.js 15", "Shows the user journey: sign in, register, verify, reports, lineage, API keys, and admin tools.", "Hasham"],
    ["Core API", "NestJS 11", "Enforces security and business rules; coordinates database, provenance, ledger, AI, and trust scoring.", "Sohaib"],
    ["Database", "MongoDB", "Stores users, roles, sessions, API keys, assets, manifests, ledger records, reports, and settings.", "Sohaib"],
    ["AI service", "FastAPI + PyTorch", "Loads the forensic model once and returns an AI-generated probability for an uploaded image.", "Shaheer"],
    ["Deployment", "Docker Compose + Caddy", "Runs services consistently and provides the production reverse-proxy/TLS boundary.", "Shared"],
], [1.05, 1.2, 3.25, 1.0])
doc.add_heading("The request path", level=2)
callout(doc, "Flow", "Browser → Next.js frontend → NestJS API → MongoDB and FastAPI AI service → trust engine → saved verification report → frontend explanation", PALE_GREEN, GREEN)
doc.add_heading("Why separate the AI service?", level=2)
bullets(doc, [
    "Python and PyTorch are the natural environment for model loading and image inference.",
    "The core platform can still return a controlled result if the AI service times out or is unavailable.",
    "A future model can be deployed without rewriting the frontend or provenance platform.",
])

# 6
new_page(doc)
page_title(doc, 5, "Hasham — Frontend Defense", "Show how the platform becomes understandable and usable.")
doc.add_heading("What Hasham should explain", level=2)
bullets(doc, [
    ("User journey: ", "signup, OTP verification, login, protected dashboard, and session-aware navigation."),
    ("Verification experience: ", "safe file selection, upload feedback, progress, errors, verdict badges, and evidence details."),
    ("Reports: ", "saved verification history, filtering, detailed reports, and JSON export."),
    ("Provenance views: ", "registered assets, version lineage, ownership, edits, and revocation state."),
    ("Developer tools: ", "API-key creation, one-time reveal, masking, rotation, revocation, limits, and usage."),
    ("Administration: ", "users, roles, platform settings, health/status information, and audit-oriented views."),
])
doc.add_heading("Suggested speaking script", level=2)
script_box(doc, "Hasham", "My responsibility was to turn complex forensic evidence into a clear user workflow. The frontend does more than upload an image: it guides the user through authentication, verification, provenance history, API access, and detailed trust reports. The interface uses the backend response directly, so the user can see why a verdict was produced rather than receiving only a label.")
doc.add_heading("What to demonstrate", level=2)
numbered(doc, [
    "Log in and show the protected dashboard.",
    "Open the verification page and explain accepted files and progress handling.",
    "Verify one image and point to the verdict, score breakdown, signals, and metadata.",
    "Open a saved report or lineage view.",
    "Briefly show API-key management or the admin area if time allows.",
])
callout(doc, "Handoff", "Now that the interface has collected the request, Sohaib will explain what the backend does with it and how the evidence is protected.", PALE_AMBER, AMBER)

# 7
new_page(doc)
page_title(doc, 6, "Sohaib — Backend Defense", "Explain how security, provenance, and evidence orchestration work.")
doc.add_heading("What Sohaib should explain", level=2)
bullets(doc, [
    ("Identity and access: ", "OTP activation, password hashing, JWT cookies, refresh-session rotation, logout, and role guards."),
    ("Provenance: ", "SHA-256, pHash/dHash, normalized metadata, signed manifests, ownership, edits, revocation, and lineage."),
    ("Ledger: ", "ordered entries, previous-entry links, Merkle roots/proofs, signed blocks, and integrity verification."),
    ("Verification orchestration: ", "metadata/hash work and the AI request run together; exact and near-match claims are resolved."),
    ("Trust engine: ", "combines independent evidence, redistributes unavailable weights, and applies hard safety caps."),
    ("Developer API: ", "peppered key hashes, scopes, rate limits, usage tracking, rotation, and revocation."),
])
doc.add_heading("Suggested speaking script", level=2)
script_box(doc, "Sohaib", "The backend is the trust boundary of ZTDPP. It never accepts a stored claim without checking it again. During verification, it recomputes hashes, validates the manifest signature, checks the ledger chain and Merkle proof, calls the AI service, and then gives the trust engine a structured evidence set. MongoDB stores the evidence and history, while uploaded verification bytes are processed in memory and discarded.")
doc.add_heading("Important wording", level=2)
callout(doc, "Say this", "The ledger is tamper-evident and platform-operated. It is a signed hash chain with Merkle proofs; we should not oversell it as a public decentralized blockchain.", PALE_RED, CORAL)
callout(doc, "Handoff", "The backend treats AI as one independent signal. Shaheer will now explain how that signal is produced and where its current limits are.", PALE_AMBER, AMBER)

# 8
new_page(doc)
page_title(doc, 7, "Shaheer — AI Defense", "Explain the model as a supporting forensic signal, with honest limitations.")
doc.add_heading("What Shaheer should explain", level=2)
bullets(doc, [
    ("Data pipeline: ", "image validation, metadata ingestion, content hashing, balanced sampling, and split checks."),
    ("Model: ", "a ConvNeXt Tiny backbone combined with a residual/high-pass stream for forensic texture cues."),
    ("Training evidence: ", "checkpoint validation accuracy 94.84%, precision 95.67%, recall 93.93%, and F1 94.79%."),
    ("Service: ", "FastAPI loads best.pt, applies the trained preprocessing, and exposes health and prediction endpoints."),
    ("Integration: ", "the backend uses timeouts, API-key protection, health probes, and fallback behavior."),
    ("Limitation: ", "a watermark-removed Gemini sample produced a severe false negative, showing a cross-generator generalization gap."),
])
doc.add_heading("Suggested speaking script", level=2)
script_box(doc, "Shaheer", "Our AI component analyzes pixel patterns that may indicate synthetic generation. The model performed strongly on its validation set, but validation accuracy is not the same as universal reliability. We found a difficult Gemini example that the model missed, so the platform reports the AI probability as evidence rather than absolute proof. Our production gate is retraining and testing on held-out generators and post-processing variations.")
doc.add_heading("If asked why the model made a mistake", level=2)
callout(doc, "Answer", "The visual fingerprints learned from the training generators may not appear in a new generator or may be weakened by editing, resizing, compression, or watermark removal. This is a data-distribution problem, so the correct response is broader training and generator-held-out testing—not simply lowering the threshold.", PALE_BLUE, BLUE)

# 9
new_page(doc)
page_title(doc, 8, "How the Trust Score Becomes a Verdict", "Keep the explanation intuitive; use the formula only if asked.")
body(doc, "The default policy gives 60% importance to provenance/metadata evidence and 40% to AI evidence. These weights can be configured. If one signal is unavailable, its weight is redistributed instead of silently treating the missing signal as a perfect result.")
table(doc, ["Evidence", "What increases confidence", "What reduces confidence"], [
    ["Provenance", "Exact or clear near match, valid signature, valid ledger proof, active claim", "No claim, ambiguous match, revoked claim, broken signature, or ledger failure"],
    ["Metadata", "Consistent source/action information and expected fields", "Missing, contradictory, or suspicious metadata and generator hints"],
    ["AI", "Low calibrated AI-generated probability", "High probability, model/evidence disagreement, or unavailable model"],
], [1.2, 2.65, 2.65])
doc.add_heading("Verdict levels", level=2)
table(doc, ["Verdict", "Easy meaning"], [
    ["Trusted", "Strong verified provenance and no critical integrity failure."],
    ["Likely authentic", "Evidence mostly supports authenticity, but confidence is not strong enough for Trusted."],
    ["Suspicious", "Important evidence is missing, weak, or conflicting."],
    ["Untrusted", "A critical integrity failure, revocation, or strong negative evidence is present."],
], [1.6, 4.9])
callout(doc, "Key point", "Hard safety caps stop a high average score from hiding a critical failure. For example, a broken signature or revoked provenance must meaningfully limit the final verdict.", PALE_RED, CORAL)

# 10
new_page(doc)
page_title(doc, 9, "Recommended Live Demo Flow", "Keep the demo controlled, short, and evidence-focused.")
table(doc, ["Time", "Presenter", "Action", "What to say"], [
    ["0:00–1:00", "Hasham", "Open dashboard and introduce the problem", "We need both trusted history and forensic evidence."],
    ["1:00–3:00", "Hasham", "Upload and verify a prepared image", "The interface validates the file and sends it securely to the API."],
    ["3:00–5:30", "Sohaib", "Explain the returned report", "These hashes, manifest checks, ledger checks, and signals were recomputed."],
    ["5:30–7:00", "Sohaib", "Show lineage or saved history", "Edits create new versions; revocation never deletes the historical record."],
    ["7:00–9:00", "Shaheer", "Show AI probability and model health", "AI is one supporting signal with a known confidence boundary."],
    ["9:00–10:30", "Team", "Show limitation and roadmap", "We know the production blockers and have explicit acceptance gates."],
    ["10:30–11:00", "Hasham", "Close", "ZTDPP explains trust instead of returning a black-box label."],
], [0.7, 0.85, 2.0, 2.95])
doc.add_heading("Prepare before entering the room", level=2)
bullets(doc, [
    "Start frontend, backend, MongoDB connection, and AI service before the defense.",
    "Keep one registered image, one edited/near-match image, and one unregistered image ready.",
    "Test the exact login account and every demo file once on the defense machine.",
    "Keep screenshots or a short screen recording as a fallback if the network or model service fails.",
    "Do not retrain, install packages, or change environment values during the live defense.",
])

# 11
new_page(doc)
page_title(doc, 10, "Be Honest About What Remains", "A strong defense separates completed engineering from future work.")
table(doc, ["Current gap", "How to explain it", "Next action"], [
    ["AI generalization", "Strong validation results, but at least one modern Gemini sample was missed.", "Broader generator corpus, held-out testing, calibration, and regression gates."],
    ["GradCAM", "The approved SRS requested a heatmap, but it is not implemented for the deployed model.", "Implement ConvNeXt GradCAM or formally approve another explanation method."],
    ["PDF forensic export", "Reports are saved and JSON can be downloaded; a formatted signed PDF is pending.", "Generate PDF from the immutable saved verification snapshot."],
    ["Performance evidence", "Builds and tests pass, but formal p50/p95/p99 and concurrency evidence is pending.", "Run reproducible load and latency benchmarks on declared hardware."],
    ["Architecture documents", "The implementation improved and changed beyond the original SRS/SDS design.", "Publish an approved SRS v3.0 and SDS v2.0 change baseline."],
    ["External C2PA", "The platform detects C2PA container presence but does not yet validate every third-party trust chain.", "Integrate a conforming verifier and explicit trust-store policy."],
], [1.35, 3.05, 2.1])
callout(doc, "Defense strategy", "Do not hide limitations. State what is implemented, show the evidence, explain why the limitation exists, and give a technically credible next step.", PALE_GREEN, GREEN)

# 12
new_page(doc)
page_title(doc, 11, "Likely Defense Questions and Short Answers", "Answer directly, then offer technical detail if the panel asks for more.")
table(doc, ["Question", "Strong short answer"], [
    ["What is the main novelty?", "Combining signed provenance, tamper-evident lineage, perceptual matching, metadata, AI evidence, and an explainable trust policy in one operational platform."],
    ["Is the ledger a blockchain?", "It is a platform-operated tamper-evident hash chain with signed blocks and Merkle proofs. It is not presented as a decentralized public blockchain."],
    ["Can AI prove an image is fake?", "No. AI provides a probability based on learned pixel patterns. ZTDPP combines it with independent provenance and integrity evidence."],
    ["Why MongoDB?", "The implemented domain contains evolving nested evidence, manifests, lineage, reports, and settings. MongoDB supports that structure and transaction-backed writes."],
    ["What happens if AI is down?", "The backend uses timeout and fallback handling, reports AI as unavailable, redistributes effective weights, and does not invent an AI result."],
    ["How is tampering detected?", "The system recomputes content hashes, manifest hashes, Ed25519 signatures, previous-entry links, Merkle proofs, and block integrity."],
    ["Do you store uploaded verification images?", "Normal verification processes bytes in memory and discards them; the saved report contains evidence, hashes, policy, and results."],
    ["What would you do before public deployment?", "Close authorization gaps, rotate production secrets, prove database and ledger recovery, run E2E/load tests, and pass generator-held-out AI gates."],
], [2.15, 4.35])
doc.add_heading("Closing statement", level=2)
script_box(doc, "Team", "ZTDPP has progressed beyond a simple AI detector into an explainable provenance platform. We can register trusted history, preserve lineage, verify integrity, integrate AI evidence, and show the user exactly why a verdict was produced. Our remaining work is clearly defined through production acceptance gates rather than an open-ended rewrite.")

doc.core_properties.title = "ZTDPP Defense Guide and Project Pitch"
doc.core_properties.subject = "Simple project explanation, presenter roles, demo flow, limitations, and defense questions"
doc.core_properties.author = "ZTDPP Project Team"
doc.settings.element.set(qn("w:updateFields"), "true")
doc.save(OUT)
print(OUT)
